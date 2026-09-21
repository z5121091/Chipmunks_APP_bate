import datetime
import os
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook

from scripts import label_sync_server
from scripts import label_templates
from scripts.label_templates import apm, boya, unpack_time


class LocalNetworkAddressTest(unittest.TestCase):
    def test_ipconfig_runs_without_a_console_window(self):
        completed_process = type('CompletedProcess', (), {
            'stdout': b'Windows IP Configuration\n',
        })()

        with patch.object(
            label_sync_server.subprocess,
            'run',
            return_value=completed_process,
        ) as run:
            label_sync_server.read_windows_ipconfig()

        _args, kwargs = run.call_args
        self.assertEqual(
            kwargs.get('creationflags'),
            getattr(label_sync_server.subprocess, 'CREATE_NO_WINDOW', 0),
        )

    def test_prefers_physical_adapter_with_default_gateway_over_vpn(self):
        output = '''
Windows IP Configuration

Ethernet adapter Ethernet:

   Connection-specific DNS Suffix  . :
   IPv4 Address. . . . . . . . . . . : 192.168.31.26
   Default Gateway . . . . . . . . . : 192.168.31.1

Tunnel adapter Company VPN:

   IPv4 Address. . . . . . . . . . . : 10.9.0.5
   Default Gateway . . . . . . . . . : 10.9.0.1
'''
        adapters = label_sync_server.parse_windows_ipconfig(output)

        self.assertEqual(
            label_sync_server.get_lan_ipv4_addresses(adapters, []),
            ['192.168.31.26'],
        )

    def test_prefers_adapter_with_default_gateway_when_two_physical_cards_exist(self):
        adapters = [
            {
                'name': 'Ethernet adapter Ethernet',
                'addresses': ['192.168.10.18'],
                'has_default_gateway': False,
            },
            {
                'name': 'Wireless LAN adapter Wi-Fi',
                'addresses': ['192.168.31.26'],
                'has_default_gateway': True,
            },
        ]

        self.assertEqual(
            label_sync_server.get_lan_ipv4_addresses(adapters, []),
            ['192.168.31.26', '192.168.10.18'],
        )

    def test_uses_environment_override_for_unusual_network_card(self):
        adapters = [{
            'name': 'Unknown adapter',
            'addresses': ['192.168.31.26'],
            'has_default_gateway': True,
        }]
        with patch.dict(os.environ, {'PALM_WAREHOUSE_SYNC_IP': '192.168.31.88'}):
            self.assertEqual(
                label_sync_server.get_lan_ipv4_addresses(adapters, []),
                ['192.168.31.88'],
            )

    def test_fallback_keeps_only_private_non_loopback_ipv4(self):
        self.assertEqual(
            label_sync_server.get_lan_ipv4_addresses([], [
                '127.0.0.1', '169.254.1.1', '198.18.0.1', '8.8.8.8', '192.168.31.26',
            ]),
            ['192.168.31.26'],
        )


class NativeLabelWithoutTraceTest(unittest.TestCase):
    def test_empty_trace_keeps_printing_without_trace_barcode(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            '型号', '标签数量', '追溯码', '批次', '生产日期', '封装',
            '存货编码', '箱号', '版本号', '标签类型', '供应商',
        ])
        sheet.append([
            'A', 960, '', 'LOT-1', '20260411', 'SOT223',
            'PN-1', 'BOX-1', '', '剩余标签', '珠海领芯科技有限公司',
        ])

        [record] = label_sync_server.get_native_label_records(workbook)
        captured_text_items = []

        def capture_text_bitmap(text_items, image_items=(), **_options):
            captured_text_items.extend(text_items)
            return b'BITMAP\r\n'

        with patch.object(
            apm,
            'build_native_text_bitmap',
            side_effect=capture_text_bitmap,
        ):
            commands = apm.build_native_apm_label(
                record,
                label_templates.NATIVE_LABEL_TEMPLATE_LEADCORE,
            )

        self.assertEqual(record['trace_no'], '')
        self.assertIn(
            (label_templates.NATIVE_LABEL_RIGHT_VALUE_X, 242, '-', False),
            captured_text_items,
        )
        self.assertNotIn(b'BARCODE 350,268', commands)
        self.assertIn(b'"A/LOT-1/SOT223/PN-1/960/20260411//BOX-1"', commands)

        record['trace_no'] = 'TRACE-1'
        with patch.object(apm, 'build_native_text_bitmap', return_value=b'BITMAP\r\n'):
            commands = apm.build_native_apm_label(
                record,
                label_templates.NATIVE_LABEL_TEMPLATE_LEADCORE,
            )

        self.assertIn(b'BARCODE 350,268', commands)
        self.assertIn(b'"A/LOT-1/SOT223/PN-1/960/20260411/TRACE-1/BOX-1"', commands)

    def test_boya_supplier_uses_its_own_layout(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            '型号', '标签数量', '追溯码', '批次', '生产日期', '封装',
            '存货编码', '箱号', '版本号', '标签类型', '供应商',
        ])
        sheet.append([
            'BY25Q128ESSIG(R)', 1255, '', 'AP5P001', '2617', 'SOP8L(208mil)',
            'BY25Q128', 'BOX-1', '', '剩余标签', '珠海博雅科技股份有限公司',
        ])

        [record] = label_sync_server.get_native_label_records(workbook)
        template = label_sync_server.get_native_label_template(record)
        captured = {}

        def capture_text_bitmap(
            text_items,
            image_items=(),
            ellipse_items=(),
            line_items=(),
        ):
            captured['text_items'] = text_items
            captured['image_items'] = image_items
            captured['ellipse_items'] = ellipse_items
            captured['line_items'] = line_items
            return b'BITMAP\r\n'

        with patch.object(
            boya,
            'build_native_text_bitmap',
            side_effect=capture_text_bitmap,
        ):
            commands = label_sync_server.build_native_label(record, template)

        self.assertEqual(template, label_templates.NATIVE_LABEL_TEMPLATE_BOYA)
        self.assertEqual(
            {
                key: record[key]
                for key in ('model', 'package', 'quantity', 'batch', 'production_date', 'source_no')
            },
            {
                'model': 'BY25Q128ESSIG(R)',
                'package': 'SOP8L(208mil)',
                'quantity': '1255',
                'batch': 'AP5P001',
                'production_date': '2617',
                'source_no': 'BOX-1',
            },
        )
        header = 'BOYA MICROELECTRONICS'
        self.assertIn(
            (
                label_templates.get_native_centered_text_x(
                    header,
                    size_pixels=label_templates.NATIVE_LABEL_BOYA_HEADER_FONT_SIZE_PIXELS,
                ),
                28,
                header,
                False,
                label_templates.NATIVE_LABEL_BOYA_HEADER_FONT_SIZE_PIXELS,
            ),
            captured['text_items'],
        )
        self.assertIn((210, 330, 'BOX-1', False), captured['text_items'])
        self.assertIn((701, 122, 'RoHS', True), captured['text_items'])
        self.assertNotIn('PASS', {item[2] for item in captured['text_items']})
        self.assertNotIn('QC03', {item[2] for item in captured['text_items']})
        self.assertEqual(len(captured['image_items']), 1)
        self.assertEqual(captured['ellipse_items'], (((694, 92, 770, 168), 3),))
        self.assertEqual(captured['line_items'], ())
        self.assertNotIn(b'BARCODE ', commands)
        self.assertIn(b'QRCODE 610,190,L,6', commands)
        self.assertIn(
            b'"BY25Q128ESSIG(R)/SOP8L(208mil)/1255/AP5P001/2617/BOX-1"',
            commands,
        )
        self.assertIn(b'PRINT 1,2', commands)


class UnpackTimeLabelTest(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        patches = patch.multiple(
            label_sync_server,
            DATA_ROOT=directory.name,
            NATIVE_PRINT_HISTORY_FILE=os.path.join(directory.name, 'history.json'),
        )
        patches.start()
        self.addCleanup(patches.stop)

    def records(self, supplier='珠海极海半导体有限公司'):
        workbook = Workbook()
        workbook.active.append(['型号', '存货编码', '标签数量', '追溯码', '标签类型', '供应商', '拆包时间'])
        for label_type, quantity in [('发货标签', 200), ('剩余标签', 2300)]:
            workbook.active.append(['DEMO', 'IC.TEST.00', quantity, '', label_type, supplier, '2026-09-15 14:30'])
        return label_sync_server.get_native_label_records(workbook)

    @patch.object(label_sync_server, 'log')
    @patch.object(label_sync_server, 'send_raw_tspl_to_printer', return_value=42)
    def test_five_labels_in_one_job_and_no_extra_prints_on_retry(self, send, _log):
        for index, supplier in enumerate(label_templates.NATIVE_LABEL_TEMPLATE_BY_SUPPLIER):
            with self.subTest(supplier=supplier):
                records = self.records(supplier)
                send.reset_mock()
                result = label_sync_server.submit_native_unpack_labels(records, f'PAIR-{index}')
                self.assertEqual(records[1]['unpacked_at'], '2026-09-15 14:30')
                self.assertEqual(result['physical_label_count'], 5)
                self.assertEqual(result['time_label_count'], 1)
                commands = send.call_args.args[1]
                self.assertEqual(commands.count(b'PRINT 1,2\r\n'), 2)
                self.assertEqual(commands.count(b'PRINT 1,1\r\n'), 1)
                self.assertTrue(commands.endswith(b'PRINT 1,1\r\n'))
                retry = label_sync_server.submit_native_unpack_labels(records, f'PAIR-{index}')
                self.assertTrue(retry['duplicate'])
                self.assertEqual(retry['physical_label_count'], 5)
                self.assertEqual(retry['time_label_count'], 1)
                send.assert_called_once()

    def test_template_keeps_exact_fields_and_source_time(self):
        record = self.records()[1]
        with patch.object(unpack_time, 'build_native_text_bitmap', return_value=b'BITMAP\r\n') as bitmap:
            commands = unpack_time.build_native_unpack_time_label(record)
        items = bitmap.call_args.args[0]
        self.assertEqual([item[2] for item in items], [
            '物料型号：', 'DEMO', '存货编码：', 'IC.TEST.00', '湿敏等级：', 'MSL-3',
            '最后拆包时间：', '2026-09-15 14:30', '状态：', '已入干燥柜 / 敞口存放',
        ])
        for x, y, text, bold, *sizes in items:
            font = unpack_time.get_unpack_time_font(bold, sizes[0] if sizes else 23)
            self.assertLessEqual(x + font.getlength(text), 800)
            self.assertLessEqual(y + font.getbbox(text, anchor='lt')[3], 400)
        self.assertIn(b'SIZE 100 mm,50 mm', commands)
        self.assertTrue(commands.endswith(b'PRINT 1,1\r\n'))

    def test_old_or_missing_time_never_becomes_reprint_time(self):
        for value in ['2026/09/15 14:30', '2026-09-15 14:30:59', datetime.datetime(2026, 9, 15, 14, 30)]:
            self.assertEqual(unpack_time.format_unpack_time(value), '2026-09-15 14:30')
        for value in [None, '', '-', 'bad timestamp']:
            self.assertEqual(unpack_time.format_unpack_time(value), '-')

    @patch.object(label_sync_server, 'log')
    @patch.object(label_sync_server, 'send_raw_tspl_to_printer')
    def test_unknown_supplier_still_skips_all_automatic_prints(self, send, _log):
        result = label_sync_server.submit_native_unpack_labels(self.records('其他供应商'), 'SKIP')
        self.assertEqual(result['physical_label_count'], 0)
        self.assertEqual(result['time_label_count'], 0)
        send.assert_not_called()

    @patch.object(label_sync_server, 'send_raw_tspl_to_printer')
    def test_template_failure_does_not_submit_a_partial_job(self, send):
        with patch.object(label_sync_server, 'build_native_unpack_time_label', side_effect=RuntimeError('font missing')):
            with self.assertRaisesRegex(RuntimeError, 'font missing'):
                label_sync_server.submit_native_unpack_labels(self.records(), 'FAILED')
        send.assert_not_called()
        self.assertEqual(label_sync_server.load_native_print_history(), {})

    @patch.object(label_sync_server, 'log')
    @patch.object(label_sync_server, 'send_raw_tspl_to_printer')
    def test_existing_four_label_history_does_not_trigger_new_prints(self, send, _log):
        label_sync_server.save_native_print_history({'OLD': {'labelCount': 2, 'copiesPerLabel': 2}})
        result = label_sync_server.submit_native_unpack_labels(self.records(), 'OLD')
        self.assertTrue(result['duplicate'])
        self.assertEqual(result['physical_label_count'], 4)
        self.assertEqual(result['time_label_count'], 0)
        send.assert_not_called()


if __name__ == '__main__':
    unittest.main()
